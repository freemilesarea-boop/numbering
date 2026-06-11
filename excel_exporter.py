"""엑셀 생성 모듈.

정규화/중복제거된 레코드 리스트를 받아 영업용 엑셀 파일을 만든다.
- 시트: 전체리스트 / 전화번호있음 / 전화번호없음 / 요약
- 헤더 bold, freeze pane, 필터, 컬럼 너비 자동 조정
- 전화번호는 문자열 유지, 지도URL은 하이퍼링크 처리
"""

from __future__ import annotations

import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from dedupe import normalize_phone

# 엑셀 컬럼 순서(계획서 12절 기준).
COLUMNS = [
    "수집일",
    "지역",
    "업종",
    "검색키워드",
    "매장명",
    "전화번호",
    "주소",
    "카테고리",
    "지도URL",
    "status",
    "영업상태",
    "메모",
    "최근연락일",
    "다음연락일",
    "담당자",
]

# 전화번호처럼 항상 문자열로 유지해야 하는 컬럼.
_TEXT_COLUMNS = {"전화번호"}
# 하이퍼링크로 처리할 컬럼.
_LINK_COLUMNS = {"지도URL"}

_HEADER_FONT = Font(bold=True)
_LINK_FONT = Font(color="0563C1", underline="single")
_MAX_COL_WIDTH = 60


def build_filename(region: str, business_type: str, output_dir: str) -> str:
    """leads_{지역}_{업종}_{YYYY-MM-DD}.xlsx 형태의 경로를 만든다."""
    today = date.today().isoformat()
    filename = f"leads_{region}_{business_type}_{today}.xlsx"
    return os.path.join(output_dir, filename)


def _write_sheet(ws, records: list[dict]) -> None:
    """워크시트에 헤더 + 데이터 행을 쓰고 스타일을 적용한다."""
    # 헤더
    ws.append(COLUMNS)
    for cell in ws[1]:
        cell.font = _HEADER_FONT

    # 데이터 행
    for record in records:
        row = []
        for col in COLUMNS:
            value = record.get(col, "")
            row.append("" if value is None else value)
        ws.append(row)

        current_row = ws.max_row
        for idx, col in enumerate(COLUMNS, start=1):
            cell = ws.cell(row=current_row, column=idx)
            if col in _TEXT_COLUMNS:
                # 전화번호 등은 숫자로 해석되지 않도록 문자열 강제
                cell.number_format = "@"
                cell.value = "" if cell.value is None else str(cell.value)
            elif col in _LINK_COLUMNS and cell.value:
                cell.hyperlink = str(cell.value)
                cell.font = _LINK_FONT

    # 헤더 고정 + 필터
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    _autosize_columns(ws)


def _autosize_columns(ws) -> None:
    """컬럼 너비를 내용 길이에 맞춰 대략적으로 조정한다."""
    for idx, col in enumerate(COLUMNS, start=1):
        letter = get_column_letter(idx)
        max_len = len(str(col))
        for cell in ws[letter]:
            if cell.value is not None:
                # 한글은 폭이 넓으므로 살짝 가중치를 준다.
                text = str(cell.value)
                length = sum(2 if ord(ch) > 127 else 1 for ch in text)
                max_len = max(max_len, length)
        ws.column_dimensions[letter].width = min(max_len + 2, _MAX_COL_WIDTH)


def _write_summary_sheet(ws, stats: dict, region: str, business_type: str) -> None:
    """요약 시트를 작성한다."""
    rows = [
        ("항목", "값"),
        ("지역", region),
        ("업종", business_type),
        ("수집일", date.today().isoformat()),
        ("원본 수집 수", stats.get("original_count", 0)),
        ("중복 제거 수", stats.get("removed_count", 0)),
        ("최종 저장 수", stats.get("final_count", 0)),
        ("전화번호 있음", stats.get("with_phone", 0)),
        ("전화번호 없음", stats.get("without_phone", 0)),
    ]
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 24
    for row in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        for cell in row:
            cell.alignment = Alignment(horizontal="left")


def export_to_excel(
    records: list[dict],
    stats: dict,
    region: str,
    business_type: str,
    output_dir: str,
) -> str:
    """레코드를 엑셀 파일로 저장하고 저장 경로를 반환한다."""
    os.makedirs(output_dir, exist_ok=True)
    path = build_filename(region, business_type, output_dir)

    with_phone = [r for r in records if normalize_phone(r.get("전화번호"))]
    without_phone = [r for r in records if not normalize_phone(r.get("전화번호"))]

    wb = Workbook()

    ws_all = wb.active
    ws_all.title = "전체리스트"
    _write_sheet(ws_all, records)

    _write_sheet(wb.create_sheet("전화번호있음"), with_phone)
    _write_sheet(wb.create_sheet("전화번호없음"), without_phone)
    _write_summary_sheet(wb.create_sheet("요약"), stats, region, business_type)

    wb.save(path)
    return path
